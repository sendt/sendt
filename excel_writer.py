import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import EXCEL_FILE, TIMEFRAMES

C_BG      = "111827"
C_TF_1M   = "1E3A5F"
C_TF_5M   = "1E4D3A"
C_TF_15M  = "4A2020"
C_TF_1H   = "3B2A50"
C_TF_1D   = "3D2B00"
C_TF_BTC  = "1A3040"
C_INFO    = "0F172A"
C_WIN     = "0F2D1A"
C_LOSS    = "2D0F0F"
C_BE      = "2D260A"
C_VAZ     = "1E1E2A"
C_ROW_E   = "161D2B"
C_ROW_O   = "111827"
C_GREEN   = "22C55E"
C_RED     = "EF4444"
C_YELLOW  = "EAB308"
C_BLUE    = "60A5FA"
C_WHITE   = "F1F5F9"
C_GRAY    = "94A3B8"

TF_LABELS = {"1m":"1 DAKİKA","5m":"5 DAKİKA","15m":"15 DAKİKA","1h":"1 SAAT","1d":"1 GÜN"}
TF_COLORS = {"1m":C_TF_1M,"5m":C_TF_5M,"15m":C_TF_15M,"1h":C_TF_1H,"1d":C_TF_1D}

INFO_COLS = ["Tarih","Saat","Gün","Seans","Yön","Sonuç"]
IND_COLS  = ["Donchian","Don%","Fiyat/Hull","RSI","RSI Önc","RSI Yön","RSI Div",
             "Stoch K","Stoch D","UT K=1","UT K=2","MACD Hist","MACD Yön","Trend",
             "Hacim","Hacim Ort","ATR%"]
IND_KEYS  = ["donchian","donchian_pct","fiyat_hull","rsi","rsi_prev","rsi_yon","rsi_div",
             "stoch_k","stoch_d","ut_bot_k1","ut_bot_k2","macd_hist","macd_yon","trend",
             "vol_lbl","vol_ratio","atr_pct"]
IND_W     = [10,7,14,6,7,8,7, 8,8,8,8,11,11,8, 9,9,7]

BTC_COLS   = ["BTC Yön"]
BTC_KEYS   = ["btc_yon"]
EXTRA_COLS = ["Giriş Fiyatı", "Çıkış Fiyatı", "Kapanış Saati", "Süre(dk)"]
ALL_COLS   = INFO_COLS + IND_COLS * len(TIMEFRAMES) + BTC_COLS + EXTRA_COLS


def _fill(c): return PatternFill("solid", fgColor=c)
def _font(bold=False, color=C_WHITE, size=9):
    return Font(bold=bold, color=color, name="Segoe UI", size=size)
def _center(): return Alignment(horizontal="center", vertical="center", wrap_text=False)
def _border():
    s = Side(style="thin", color="1F2937")
    return Border(left=s, right=s, top=s, bottom=s)


def _create_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TRADE LOG"
    last = get_column_letter(len(ALL_COLS))

    # Satır 1 — başlık
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = "WLD/USDT — TRADE JOURNAL"
    ws["A1"].font = Font(bold=True, color=C_WHITE, name="Segoe UI", size=13)
    ws["A1"].fill = _fill(C_BG)
    ws["A1"].alignment = _center()

    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = "Paper Trading · 31x Kaldıraç · Hedef: 50 İşlem · Bot Geliştirme Verisi"
    ws["A2"].font = _font(color=C_GRAY)
    ws["A2"].fill = _fill(C_BG)
    ws["A2"].alignment = _center()

    # Satır 3 — bölüm başlıkları
    n_info = len(INFO_COLS)
    ws.merge_cells(f"A3:{get_column_letter(n_info)}3")
    ws["A3"] = "📋 İŞLEM"
    ws["A3"].font = _font(bold=True)
    ws["A3"].fill = _fill(C_INFO)
    ws["A3"].alignment = _center()

    col = n_info + 1
    for tf in TIMEFRAMES:
        end = col + len(IND_COLS) - 1
        sl, el = get_column_letter(col), get_column_letter(end)
        ws.merge_cells(f"{sl}3:{el}3")
        cell = ws[f"{sl}3"]
        cell.value = f"⏱ {TF_LABELS[tf]}"
        cell.font = _font(bold=True)
        cell.fill = _fill(TF_COLORS[tf])
        cell.alignment = _center()
        col = end + 1

    # BTC bölümü
    btc_col = get_column_letter(col)
    ws[f"{btc_col}3"] = "BTC"
    ws[f"{btc_col}3"].font = _font(bold=True)
    ws[f"{btc_col}3"].fill = _fill(C_TF_BTC)
    ws[f"{btc_col}3"].alignment = _center()

    # EXTRA bölümü (Çıkış Fiyatı vb.)
    extra_start = col + 1
    for j, lbl in enumerate(EXTRA_COLS):
        ec = get_column_letter(extra_start + j)
        ws[f"{ec}3"] = lbl
        ws[f"{ec}3"].font = _font(bold=True)
        ws[f"{ec}3"].fill = _fill(C_INFO)
        ws[f"{ec}3"].alignment = _center()

    # Satır 4 — sütun isimleri
    n_btc_extra = len(BTC_COLS) + len(EXTRA_COLS)
    for i, h in enumerate(ALL_COLS, start=1):
        cell = ws.cell(row=4, column=i, value=h)
        if i <= n_info:
            cell.fill = _fill(C_INFO)
        elif i <= n_info + len(IND_COLS) * len(TIMEFRAMES):
            tf_idx = min((i - n_info - 1) // len(IND_COLS), len(TIMEFRAMES) - 1)
            cell.fill = _fill(TF_COLORS[TIMEFRAMES[tf_idx]])
        elif i <= n_info + len(IND_COLS) * len(TIMEFRAMES) + len(BTC_COLS):
            cell.fill = _fill(C_TF_BTC)
        else:
            cell.fill = _fill(C_INFO)
        cell.font = _font(bold=True, size=8)
        cell.alignment = _center()
        cell.border = _border()

    # Satır yükseklikleri
    for r, h in [(1,22),(2,14),(3,16),(4,14)]:
        ws.row_dimensions[r].height = h

    # Sütun genişlikleri
    for i in range(1, len(ALL_COLS) + 1):
        letter = get_column_letter(i)
        if i <= n_info:
            widths = [10, 7, 7, 9, 7, 10]
            ws.column_dimensions[letter].width = widths[i - 1]
        elif i <= n_info + len(IND_COLS) * len(TIMEFRAMES):
            ws.column_dimensions[letter].width = IND_W[(i - n_info - 1) % len(IND_COLS)]
        else:
            ws.column_dimensions[letter].width = 9

    ws.freeze_panes = "A5"
    _legend(wb)
    return wb


def _legend(wb):
    ws = wb.create_sheet("AÇIKLAMALAR")
    rows = [
        ("Don%",), (None,"0%","Alt bantta"),  (None,"50%","Ortada"), (None,"100%","Üst bantta"),
        (),
        ("MACD Yön",),
        (None,"Y.ARTAN","Yeşil hist büyüyor"), (None,"Y.AZALAN","Yeşil hist küçülüyor"),
        (None,"K.ARTAN","Kırmızı hist büyüyor"), (None,"K.AZALAN","Kırmızı hist küçülüyor"),
        (),
        ("RSI Div",),
        (None,"Bull","Fiyat düşük dip, RSI yüksek dip → yukarı dönüş"),
        (None,"Bear","Fiyat yüksek tepe, RSI düşük tepe → aşağı dönüş"),
        (),
        ("Seans (UTC)",),
        (None,"Asya","00:00-08:00"), (None,"Avrupa","08:00-15:00"),
        (None,"ABD","15:00-22:00"),  (None,"Sakin","22:00-00:00"),
    ]
    for r, data in enumerate(rows, start=1):
        for c, val in enumerate(data, start=1):
            if val: ws.cell(row=r, column=c, value=val)


def _val_color(col, val):
    if val in ("YEŞİL","BUY","WIN","LONG","YUKARI","Y.ARTAN","Bull","Asya","YÜKSEK"):
        return C_GREEN
    if val in ("KIRMIZI","SELL","LOSS","SHORT","AŞAĞI","K.ARTAN","Bear","DÜŞÜK"):
        return C_RED
    if val in ("BE","MID","Y.AZALAN","K.AZALAN","YATAY","-","NORMAL"):
        return C_YELLOW
    if col in ("RSI","RSI Önc","Stoch K","Stoch D") and val not in (None,"-"):
        try:
            v = float(val)
            if v > 70: return C_RED
            if v < 30: return C_GREEN
            return C_WHITE
        except (ValueError, TypeError):
            pass
    if col == "Don%" and val not in (None,"-"):
        try:
            v = float(val)
            if v > 80: return C_RED
            if v < 20: return C_GREEN
            return C_WHITE
        except (ValueError, TypeError):
            pass
    if col == "MACD Hist" and val not in (None,"-"):
        try:
            return C_GREEN if float(val) >= 0 else C_RED
        except (ValueError, TypeError):
            pass
    return None


def append_row(trade: dict):
    os.makedirs(os.path.dirname(EXCEL_FILE), exist_ok=True)
    if os.path.isfile(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb["TRADE LOG"]
    else:
        wb = _create_workbook()
        ws = wb["TRADE LOG"]

    result = trade.get("result","").upper()
    row_color = {"WIN":C_WIN,"LOSS":C_LOSS,"BE":C_BE,"VAZGEÇTİM":C_VAZ}.get(
        result, C_ROW_E if ws.max_row % 2 == 0 else C_ROW_O)
    row_fill = _fill(row_color)
    next_row = ws.max_row + 1

    values = [
        trade.get("date",""), trade.get("time",""),
        trade.get("gun",""), trade.get("seans",""),
        trade.get("direction",""), result,
    ]
    for tf in TIMEFRAMES:
        tf_data = trade.get("indicators",{}).get(tf,{})
        for k in IND_KEYS:
            values.append(tf_data.get(k,"-"))
    values.append(trade.get("btc_yon","-"))
    values.append(trade.get("entry_price", "-"))
    values.append(trade.get("close_price", "-"))
    values.append(trade.get("close_time",  "-"))
    values.append(trade.get("duration_min","-"))

    for col_i, (val, col_name) in enumerate(zip(values, ALL_COLS), start=1):
        cell = ws.cell(row=next_row, column=col_i, value=val)
        cell.fill  = row_fill
        cell.alignment = _center()
        cell.border = _border()
        cell.font  = _font(color=_val_color(col_name, val) or C_WHITE)

    ws.row_dimensions[next_row].height = 14
    wb.save(EXCEL_FILE)
